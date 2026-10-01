"""Per-site checklist workbooks with an English copy of every sheet.

For each site's working Excel file (source language es / fr / it / pt) this writes a copy to
data/site_checklists/<SITE_ID>_<original file name>, adding after every sheet an identical sheet "<name> (EN)"
— same layout, merged cells, column widths, fills and fonts — with every text cell translated into English
(data/translations_en.json). Requirement IDs, dates and numbers are left as they are.
Greece's checklist is already in English and is copied unchanged.

Usage:  python data_prep/build_site_workbooks.py <folder with the source workbooks> [<second folder> ...]
        (the Excel working files and the Chetumal / Altamira files from the Amara NZero site-visit folder)
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "site_checklists"
TR = json.loads((ROOT / "data" / "translations_en.json").read_text(encoding="utf-8"))

SITE_FILES = {  # site_id -> source file name, source language
    "ES-MAD": ("2026_06_Checklist_Oficina_Madrid.xlsx", "es"),
    "ES-MEC": ("2026_06_Checklist_Meco_Warehouse.xlsx", "es"),
    "ES-SEV": ("2026_06_Checklist_Sevilla_Warehouse.xlsx", "es"),
    "ES-VAL": ("2026_06_Checklist_Valencia_Warehouse.xlsx", "es"),
    "ES-AND": ("2026_06_Checklist_Andoain_Warehouse.xlsx", "es"),
    "ES-BUR": ("2026_06_Checklist_Burgos_Warehouse.xlsx", "es"),
    "ES-VIT": ("2026-06_Checklist_ES_Planta_Industrial_Vitoria.xlsx", "es"),
    "ES-ACO": ("2026-06_Checklist_ACoruna_Amara_NZero_Repsol_Requisitos_legales.xlsx", "es"),
    "PT-AVE": ("2026_06_Amara_NZero_Office_Aveiro_Checklist.xlsx", "pt"),
    "MX-MTY": ("2026_06_Checklist_MX_Almacen_Monterrey.xlsx", "es"),
    "MX-CDMX": ("2026_06_Checklist_Oficina_CDMX_Mexico.xlsx", "es"),
    "MX-INO": ("2026_06_Site_Visit_Report_EPC_INOAC.xlsx", "es"),
    "MX-IBE": ("2026_06_Checklist_Almacen_Cliente_Mexico_Iberdrola.xlsx", "es"),
    "MX-BRA": ("2026_06_Checklist_Almacen_Cliente Braskem_Mexico.xlsx", "es"),
    "MX-SCH": ("2026_06_Site_Visit_Report_Schneider_Servicios_generales_Monterrey.xlsx", "es"),
    "MX-CHE": ("2026_06_Site_Visit_Report_Chetumal_Servicios_generales.xlsx", "es"),
    "MX-ALT": ("2026_06_Checklist_Almacen y servicios_Cliente_Mexico_Altamira_Quantum.xlsx", "es"),
    "CO-HOM": ("2026_06_Checklist_Colombia_Teletrabajo.xlsx", "es"),
    "CO-EPC": ("2026_06_Colombia_EPC_Solar_Checklist_AmaraNZero.xlsx", "es"),
    "FR-OFF": ("2026_06_Checklist_Amara_NZero_France_Office_FR.xlsx", "fr"),
    "FR-EPC": ("2026_06_Checklist_Amara_NZero_France_EPC.xlsx", "fr"),
    "GR-ATH": ("2026_06_Amara_NZero_Checklist_Greece.xlsx", "en"),
    "IT-BUC": ("2026_06_Checklist_IT_Office_Buccinasco.xlsx", "it"),
}
ID_RE = re.compile(r"[A-Z]{2,4}(-[A-Z0-9]+)+")


def translate(v):
    if (not isinstance(v, str) or v.startswith("=") or not re.search(r"[A-Za-zÀ-ÿ]{3,}", v)
            or ID_RE.fullmatch(v.strip())):
        return v, True
    t = TR.get(v) or TR.get(v.strip())
    return (t, True) if t else (v, False)


def last_used_row(ws):
    """Real last row (Sevilla's sheet reports ~1M rows of empty formatting)."""
    last = 0
    for r, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(c not in (None, "") for c in row):
            last = r
        elif r - last > 300:
            break
    return last


def build(src_dirs):
    OUT.mkdir(parents=True, exist_ok=True)
    report = {}
    for sid, (fname, lang) in SITE_FILES.items():
        path = next((Path(d) / fname for d in src_dirs if (Path(d) / fname).exists()), None)
        if path is None:
            print("MISSING", sid, fname)
            continue
        wb = openpyxl.load_workbook(path)
        untranslated = []
        if lang != "en":
            for ws in list(wb.worksheets):
                last = last_used_row(ws)
                if ws.max_row > last + 50:  # drop trailing empty formatted rows so the copy stays small
                    ws.delete_rows(last + 1, ws.max_row - last)
                en = wb.copy_worksheet(ws)
                en.title = (ws.title[:26] + " (EN)")
                wb.move_sheet(en, offset=wb.index(ws) + 1 - wb.index(en))
                en.sheet_view.zoomScale = ws.sheet_view.zoomScale
                en.freeze_panes = ws.freeze_panes
                for row in en.iter_rows():
                    for c in row:
                        if isinstance(c.value, str):
                            new, ok = translate(c.value)
                            if not ok:
                                untranslated.append(f"{en.title}!{c.coordinate}: {c.value[:60]}")
                            c.value = new
        out = OUT / f"{sid}_{fname}"
        wb.save(out)
        report[sid] = {"file": out.name, "sheets": wb.sheetnames, "untranslated_cells": untranslated}
        print(f"{sid:8s} {len(wb.sheetnames)} sheets  untranslated={len(untranslated)}")
    (OUT / "_build_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    build(sys.argv[1:])
