"""Excel exports built from the same data the dashboard uses.

build_records_workbook()  -> complete records: legal requirements (page 1) and the assessed checklist / site-visit
                             report findings (page 2), each in an Original-language sheet and an English sheet.
build_graph_workbook()    -> one sheet per dashboard chart: the table behind it (with SUM formulas for checking)
                             and a native Excel chart, plus QA reconciliation sheets.
Both return bytes, so the app can offer them as downloads and data_prep/export_excel.py can write them to disk.
"""
import io
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
DATA = ROOT / "data"

NAVY, GREEN, ORANGE, RED, GREY = "1B1F3B", "1FA276", "F0A02C", "E24B47", "F2F3F5"
STATUS = ["Compliant", "Partially compliant", "Non-compliant"]
CRIT = ["High", "Medium", "Low"]
TYPES = ["Office", "EPC", "Factory", "Warehouse", "Services"]
COUNTRIES = ["Spain", "Portugal", "Mexico", "Colombia", "France", "Greece", "Italy"]
CLUSTERS = ["Spain & Portugal", "Mexico & Colombia", "France", "Greece & Italy"]
LANG_NAME = {"es": "Spanish", "fr": "French", "it": "Italian", "pt": "Portuguese", "en": "English"}

HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
NOTE_FONT = Font(italic=True, size=9, color="4A4F66")
THIN = Side(style="thin", color="D5D8E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _load():
    rows = json.loads((DATA / "checklist.json").read_text(encoding="utf-8"))
    tr = json.loads((DATA / "translations_en.json").read_text(encoding="utf-8"))
    deck = json.loads((DATA / "deck_results.json").read_text(encoding="utf-8"))
    reports = json.loads((DATA / "site_reports.json").read_text(encoding="utf-8"))
    qa_path = DATA / "qa_results.json"
    qa = json.loads(qa_path.read_text(encoding="utf-8")) if qa_path.exists() else {}
    return rows, tr, deck, reports, qa


def _en(tr, v):
    return tr.get(v, tr.get(v.strip(), v)) if isinstance(v, str) and v else v


def _table(ws, top, left, headers, data, widths=None, wrap_cols=()):
    """Write a formatted table; returns (first_data_row, last_data_row)."""
    for j, h in enumerate(headers):
        c = ws.cell(top, left + j, h)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(data, 1):
        for j, v in enumerate(row):
            c = ws.cell(top + i, left + j, v)
            c.border = BORDER
            if j in wrap_cols:
                c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(left + j)].width = w
    return top + 1, top + len(data)


def _title(ws, text, note=""):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = NOTE_FONT


def _stack_chart(ws, title, cats_ref, first_col, n_series, min_row, max_row, anchor, colors, y_title="Requirements",
                 horizontal=False, pct=False, width=22, height=11, stacked=True, titles=None):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = ("percentStacked" if pct else "stacked") if stacked else "clustered"
    if stacked:
        ch.overlap = 100
    ch.title = title
    ch.y_axis.title = y_title
    ch.height, ch.width = height, width
    if titles:  # explicit series names (blocks that do not sit directly under the header row)
        data = Reference(ws, min_col=first_col, max_col=first_col + n_series - 1, min_row=min_row, max_row=max_row)
        ch.add_data(data, titles_from_data=False)
        for ser, t in zip(ch.series, titles):
            ser.tx = SeriesLabel(v=t)
    else:
        data = Reference(ws, min_col=first_col, max_col=first_col + n_series - 1, min_row=min_row - 1, max_row=max_row)
        ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats_ref)
    for s, col in zip(ch.series, colors):
        s.graphicalProperties.solidFill = col
        s.graphicalProperties.line.solidFill = "FFFFFF"
    ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False, showLeaderLines=False)
    ch.dataLabels.numFmt = '0;-0;;@'
    ch.dataLabels.showVal = True
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)
    return ch


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Workbook 1 — complete records
# ══════════════════════════════════════════════════════════════════════════════
REQ_COLS = [("country", "Country"), ("cluster", "Region (cluster)"), ("site_id", "Site ID"), ("site", "Site"),
            ("business_type", "Type of business"), ("lang", "Source language"), ("id", "Requirement ID"),
            ("installation", "Installation"), ("facility_type", "Installation type"), ("activity", "Key activity"),
            ("norm", "Standard / Reference"), ("eu_directive", "EU Directive / Regulation"),
            ("requirement", "Legal requirement"), ("question", "Assessment question"),
            ("evidence", "Required evidence"), ("notes", "Notes"), ("documents", "Documents or certifications required"),
            ("frequency", "Frequency / Legal deadline")]
ASSESS_COLS = [("criticality", "Criticality level (as recorded)"), ("criticality_level", "Criticality (normalised)"),
               ("status", "Legal compliance status (as recorded)"), ("status_source", "Status source"),
               ("status_level", "Status (normalised)"), ("reason", "Reason for status"),
               ("amara_comments", "Amara team comments"), ("missing_document", "Missing document"),
               ("responsible", "Responsible")]
TEXT_FIELDS = {"installation", "facility_type", "activity", "norm", "eu_directive", "requirement", "question", "evidence",
               "notes", "documents", "frequency", "criticality", "status", "reason", "amara_comments",
               "missing_document", "responsible"}


def _sheet_rows(rows, cols, tr, english):
    out = []
    for r in rows:
        line = []
        for f, _ in cols:
            v = r.get(f, "")
            if f == "lang":
                v = LANG_NAME.get(v, v)
            elif english and f in TEXT_FIELDS:
                v = _en(tr, v)
            line.append(v)
        out.append(line)
    return out


def build_records_workbook():
    rows, tr, deck, reports, qa = _load()
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    _title(ws, "Amara NZero · Legal compliance records (complete, unfiltered)",
           "dss+ Safety, Legal Compliance & Culture Assessment · June 2026 · Confidential")
    info = [
        ("1 Requirements (Original)", "Page 1 · every legal requirement identified per country / site / type of business "
                                      "(no criticality, no results) in the language of the source checklist."),
        ("1 Requirements (English)", "Same records, English translation."),
        ("2 Checklist (Original)", "Page 2 · the same requirements with the site assessment: criticality, compliance status, "
                                   "reason, comments — as recorded in the working Excel files."),
        ("2 Checklist (English)", "Same records, English translation."),
        ("2 Report activities (Orig)", "Page 2 · site-visit report 'Legal compliance assessment': key activity, regulations, "
                                       "compliance dot, comments — in the report's language."),
        ("2 Report activities (Eng)", "Same, English."),
        ("2 Report findings (Orig)", "Page 2 · strengths, areas for improvement, quick wins and recommendations per site."),
        ("2 Report findings (Eng)", "Same, English."),
        ("Status mapping", "How each recorded status value maps to the normalised status used in charts."),
        ("Criticality mapping", "How each recorded criticality value maps to High / Medium / Low."),
        ("Sites", "The 23 sites: country, cluster, type of business, source file, record count."),
    ]
    _table(ws, 4, 1, ["Sheet", "Content"], info, widths=[30, 120], wrap_cols=(1,))
    ws.cell(17, 1, "Status source").font = Font(bold=True)
    ws.cell(18, 1, "text = status written in the checklist cell; cell colour = status cell empty, status read from its "
                   "fill colour (France checklists; one Burgos row).").font = NOTE_FONT
    ws.cell(19, 1, "Pendiente / Pendente / Pending (evidence pending) is normalised to Non-compliant.").font = NOTE_FONT

    req_cols = REQ_COLS
    full_cols = REQ_COLS + ASSESS_COLS
    widths_req = [11, 16, 9, 26, 12, 11, 13, 20, 16, 30, 40, 26, 60, 50, 45, 45, 35, 35]
    widths_full = widths_req + [14, 14, 16, 12, 18, 50, 40, 25, 18]
    wrap_req = tuple(range(9, len(req_cols)))
    for name, cols, eng, widths in [("1 Requirements (Original)", req_cols, False, widths_req),
                                    ("1 Requirements (English)", req_cols, True, widths_req),
                                    ("2 Checklist (Original)", full_cols, False, widths_full),
                                    ("2 Checklist (English)", full_cols, True, widths_full)]:
        s = wb.create_sheet(name)
        data = _sheet_rows(rows, cols, tr, eng)
        _table(s, 1, 1, [h for _, h in cols], data, widths=widths, wrap_cols=wrap_req)
        s.freeze_panes = "E2"
        s.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data) + 1}"

    site_meta = {}
    for r in rows:
        site_meta.setdefault(r["site_id"], (r["country"], r["cluster"], r["site"], r["business_type"], r["lang"]))
    for name, key in [("2 Report activities (Orig)", "orig"), ("2 Report activities (Eng)", "en")]:
        s = wb.create_sheet(name)
        data = []
        for sid, rep in reports.items():
            c, cl, site, bt, _ = site_meta.get(sid, ("", "", sid, "", ""))
            for n, ka in enumerate(rep.get("key_activities", []), 1):
                act = ka["activity"].get(key) or ka["activity"].get("en")
                com = "\n".join("• " + (x.get(key) or x.get("en")) for x in ka.get("comments", []))
                data.append([c, cl, sid, site, bt, rep.get("report_file", ""), LANG_NAME.get(rep.get("report_language"), ""),
                             rep.get("visit_date", ""), n, act, "\n".join(ka.get("regulations", [])), ka.get("status", ""), com])
        _table(s, 1, 1, ["Country", "Region (cluster)", "Site ID", "Site", "Type of business", "Report file",
                         "Report language", "Visit date", "#", "Key activity", "Regulations", "Compliance (dot)",
                         "Comments / missing evidence"], data,
               widths=[11, 16, 9, 26, 12, 34, 11, 18, 5, 34, 30, 18, 90], wrap_cols=(9, 10, 12))
        s.freeze_panes = "E2"
        s.auto_filter.ref = f"A1:M{len(data) + 1}"
    for name, key in [("2 Report findings (Orig)", "orig"), ("2 Report findings (Eng)", "en")]:
        s = wb.create_sheet(name)
        data = []
        for sid, rep in reports.items():
            c, cl, site, bt, _ = site_meta.get(sid, ("", "", sid, "", ""))
            for sec, label in [("summary_strengths", "Strength"), ("summary_improvements", "Area for improvement"),
                               ("quick_wins", "Quick win"), ("recommendations", "Recommendation")]:
                for n, x in enumerate(rep.get(sec, []), 1):
                    data.append([c, cl, sid, site, bt, label, n, x.get(key) or x.get("en")])
        _table(s, 1, 1, ["Country", "Region (cluster)", "Site ID", "Site", "Type of business", "Section", "#", "Text"],
               data, widths=[11, 16, 9, 26, 12, 20, 5, 120], wrap_cols=(7,))
        s.freeze_panes = "E2"
        s.auto_filter.ref = f"A1:H{len(data) + 1}"

    s = wb.create_sheet("Status mapping")
    cnt = Counter((LANG_NAME.get(r["lang"]), r["status"] or "(empty cell)", r["status_source"], r["status_level"]) for r in rows)
    _table(s, 1, 1, ["Source language", "Status as recorded", "Status source", "Normalised status", "Records"],
           [list(k) + [v] for k, v in sorted(cnt.items())], widths=[16, 28, 14, 22, 10])
    s = wb.create_sheet("Criticality mapping")
    cnt = Counter((r["criticality"] or "(empty)", r["criticality_level"]) for r in rows)
    _table(s, 1, 1, ["Criticality as recorded", "Normalised", "Records"], [list(k) + [v] for k, v in sorted(cnt.items())],
           widths=[26, 16, 10])
    s = wb.create_sheet("Sites")
    cnt = Counter(r["site_id"] for r in rows)
    data = [[sid, *site_meta[sid][:4], LANG_NAME.get(site_meta[sid][4]), cnt[sid]] for sid in site_meta]
    _table(s, 1, 1, ["Site ID", "Country", "Region (cluster)", "Site", "Type of business", "Checklist language", "Records"],
           data, widths=[9, 11, 18, 36, 14, 16, 9])
    n = len(data) + 1
    s.cell(n + 1, 6, "Total").font = Font(bold=True)
    s.cell(n + 1, 7, f"=SUM(G2:G{n})").font = Font(bold=True)
    return _bytes(wb)


# ══════════════════════════════════════════════════════════════════════════════
# Workbook 2 — data behind every chart (+ QA)
# ══════════════════════════════════════════════════════════════════════════════
def _split_regs(text):
    parts = re.split(r"\s\+\s|\+|;|\n|\s—\s|\s–\s(?=[A-Z])", text or "")
    out = []
    for p in parts:
        p = re.sub(r"\s*\(.*?\)\s*", " ", p).strip(" .,:-–—")
        p = re.sub(r"\s+", " ", p)
        if 3 <= len(p) <= 60 and re.search(r"\d", p):
            out.append(p)
    return out


def _counts(rows):
    return {c: [sum(1 for r in rows if r["criticality_level"] == c and r["status_level"] == s) for s in STATUS] for c in CRIT}


def _crit_sheet(wb, name, title, note, blocks, chart_per_block=True, cols_per_row=2):
    """blocks: list of (label, {crit: [c,p,nc]}). One table + one stacked chart per block."""
    s = wb.create_sheet(name)
    _title(s, title, note)
    top = 4
    for bi, (label, fig) in enumerate(blocks):
        s.cell(top, 1, label).font = Font(bold=True, color=NAVY)
        data = [[c, *fig[c], f"=SUM(B{top + 2 + i}:D{top + 2 + i})"] for i, c in enumerate(CRIT)]
        r0, r1 = _table(s, top + 1, 1, ["Criticality", *STATUS, "Total"], data, widths=[16, 12, 18, 14, 9])
        s.cell(r1 + 1, 1, "Total").font = Font(bold=True)
        for j, col in enumerate("BCDE"):
            s.cell(r1 + 1, 2 + j, f"=SUM({col}{r0}:{col}{r1})").font = Font(bold=True)
        if chart_per_block:
            _stack_chart(s, label, Reference(s, min_col=1, min_row=r0, max_row=r1), 2, 3, r0, r1,
                         f"{'G' if bi % cols_per_row == 0 else 'Q'}{top}", [GREEN, ORANGE, RED], width=15, height=7.5)
        top = r1 + 12 if chart_per_block else r1 + 3
    return s


def build_graph_workbook():
    rows, tr, deck, reports, qa = _load()
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    _title(ws, "Amara NZero · Data behind every dashboard chart",
           "Each sheet = one chart: the table it plots (totals as SUM formulas you can check) and the same chart in Excel.")
    contents = [
        ("P1 Country x type", "Page 1 · requirements by country and type of business (treemap + matrix)"),
        ("P1 Regulations", "Page 1 · most-referenced regulations (top 15 charted, full list below)"),
        ("P1 Requirements per site", "Page 1 · requirements per site"),
        ("P1 KPIs", "Page 1 · headline numbers"),
        ("P2 Site criticality", "Page 2 · 'Compliance status by criticality' for every site (deck figures, or checklist where the deck has no site chart)"),
        ("P2 Site status (checklist)", "Page 2 · overall status of all checklist items per site (donut) + KPIs"),
        ("P2 Crit x status (checklist)", "Page 2 · checklist records by criticality and status per site (table under the donut)"),
        ("P2 Key activity (checklist)", "Page 2 · compliance by key activity per site"),
        ("P2 Report dots", "Page 2 · site-visit report compliance dots per site"),
        ("P3 Cluster business", "Page 3 · compliance status per criticality of business, per cluster (deck)"),
        ("P3 Cluster KPIs", "Page 3 · cluster KPI tiles"),
        ("P3 Cluster mix", "Page 3 · compliance mix by type of business (%)"),
        ("P3 Site heatmap", "Page 3 · share compliant by site and criticality"),
        ("P4 Global high", "Page 4 · high criticality by type of business and cluster (deck)"),
        ("P4 High by cluster", "Page 4 · one chart per cluster, high criticality"),
        ("P4 High mix", "Page 4 · high-criticality mix per cluster and per type of business (%)"),
        ("P4 All criticality", "Page 4 · all criticality levels per cluster"),
        ("P4 KPIs", "Page 4 · headline numbers"),
        ("QA Deck vs checklist", "QA · deck chart figures compared with counts from the checklists"),
        ("QA Deck internal", "QA · deck charts that do not equal the sum of their own site charts"),
        ("QA Reconciliation", "QA · source Excel → dataset reconciliation and cross-checks"),
    ]
    _table(ws, 4, 1, ["Sheet", "Chart / content"], contents, widths=[30, 120], wrap_cols=(1,))

    # ── Page 1 ──
    s = wb.create_sheet("P1 Country x type")
    _title(s, "Requirements by country and type of business", "Source: 23 site checklists (1,002 requirements).")
    data = []
    for i, c in enumerate(COUNTRIES):
        r = [c] + [sum(1 for x in rows if x["country"] == c and x["business_type"] == t) for t in TYPES]
        data.append(r + [f"=SUM(B{5 + i}:F{5 + i})"])
    r0, r1 = _table(s, 4, 1, ["Country", *TYPES, "Total"], data, widths=[14, 10, 10, 10, 11, 10, 10])
    s.cell(r1 + 1, 1, "Total").font = Font(bold=True)
    for j in range(6):
        col = get_column_letter(2 + j)
        s.cell(r1 + 1, 2 + j, f"=SUM({col}{r0}:{col}{r1})").font = Font(bold=True)
    s.cell(r1 + 2, 1, f"Check: equals total records ({len(rows)})").font = NOTE_FONT
    _stack_chart(s, "Requirements by country and type of business", Reference(s, min_col=1, min_row=r0, max_row=r1),
                 2, 5, r0, r1, "J4", ["2F6DB5", "8A8FA3", NAVY, GREEN, "C4A11A"])

    s = wb.create_sheet("P1 Regulations")
    _title(s, "Most-referenced regulations", "Regulation references split from the 'Standard / Reference' column; count = "
                                             "number of requirements citing each.")
    cnt = Counter(x for r in rows for x in set(_split_regs(r["norm"])))
    top = cnt.most_common()
    r0, r1 = _table(s, 4, 1, ["Regulation", "Requirements citing it"], top, widths=[40, 12])
    _stack_chart(s, "Top 15 regulations", Reference(s, min_col=1, min_row=r0, max_row=r0 + 14), 2, 1, r0, r0 + 14,
                 "E4", [NAVY], horizontal=True, stacked=False, height=12)

    s = wb.create_sheet("P1 Requirements per site")
    _title(s, "Requirements per site")
    per = Counter(r["site_id"] for r in rows)
    meta = {}
    for r in rows:
        meta.setdefault(r["site_id"], r)
    order = sorted(meta, key=lambda k: (COUNTRIES.index(meta[k]["country"]), -per[k]))
    data = [[meta[k]["country"], k, meta[k]["site"], meta[k]["business_type"], per[k]] for k in order]
    r0, r1 = _table(s, 4, 1, ["Country", "Site ID", "Site", "Type of business", "Requirements"], data,
                    widths=[11, 9, 38, 14, 12])
    s.cell(r1 + 1, 4, "Total").font = Font(bold=True)
    s.cell(r1 + 1, 5, f"=SUM(E{r0}:E{r1})").font = Font(bold=True)
    _stack_chart(s, "Requirements per site", Reference(s, min_col=3, min_row=r0, max_row=r1), 5, 1, r0, r1, "H4",
                 [GREEN], horizontal=True, stacked=False, height=14)

    s = wb.create_sheet("P1 KPIs")
    _title(s, "Page 1 headline numbers")
    kp = [("Legal requirements", len(rows)), ("Key activities (distinct, as written)", len({r["activity"] for r in rows})),
          ("Regulations referenced (distinct)", len(cnt)), ("Sites", len(meta)),
          ("Countries", len({r["country"] for r in rows})), ("Types of business", len({r["business_type"] for r in rows}))]
    _table(s, 3, 1, ["Measure", "Value"], kp, widths=[40, 10])

    # ── Page 2 ──
    blocks, src = [], []
    by = {}
    for r in rows:
        by.setdefault(r["site_id"], []).append(r)
    for sid in order:
        if sid in deck["site_figures"]:
            blocks.append((f"{sid} · {meta[sid]['site']} (deck)", deck["site_figures"][sid]))
        else:
            blocks.append((f"{sid} · {meta[sid]['site']} (checklist)", _counts(by[sid])))
    _crit_sheet(wb, "P2 Site criticality", "Compliance status by criticality · per site",
                "Deck = figure as printed on the cluster-deck slide. Checklist = counted from the site checklist "
                "(site has no own chart in the deck).", blocks)

    s = wb.create_sheet("P2 Site status (checklist)")
    _title(s, "Overall status of all checklist items per site", "Counts of normalised status from the checklists.")
    cols = ["Compliant", "Partially compliant", "Non-compliant", "Not applicable", "Not assessed"]
    data = []
    for i, sid in enumerate(order):
        c = Counter(r["status_level"] for r in by[sid])
        hi = sum(1 for r in by[sid] if r["criticality_level"] == "High" and r["status_level"] in ("Non-compliant", "Partially compliant"))
        data.append([sid, meta[sid]["site"], *[c.get(k, 0) for k in cols], f"=SUM(C{5 + i}:G{5 + i})", hi])
    r0, r1 = _table(s, 4, 1, ["Site ID", "Site", *cols, "Total", "High-criticality gaps"], data,
                    widths=[9, 36, 11, 12, 13, 12, 11, 9, 14])
    s.cell(r1 + 1, 2, "Total").font = Font(bold=True)
    for j in range(7):
        col = get_column_letter(3 + j)
        s.cell(r1 + 1, 3 + j, f"=SUM({col}{r0}:{col}{r1})").font = Font(bold=True)
    _stack_chart(s, "Checklist status per site", Reference(s, min_col=2, min_row=r0, max_row=r1), 3, 5, r0, r1, "L4",
                 [GREEN, ORANGE, RED, "B8BCC6", "D9DCE2"], horizontal=True, height=15)

    s = wb.create_sheet("P2 Crit x status (checklist)")
    _title(s, "Checklist records by criticality and status · per site")
    data = []
    for sid in order:
        for c in CRIT + ["Not rated"]:
            vals = [sum(1 for r in by[sid] if r["criticality_level"] == c and r["status_level"] == k) for k in cols]
            if sum(vals):
                data.append([sid, meta[sid]["site"], c, *vals])
    r0, r1 = _table(s, 4, 1, ["Site ID", "Site", "Criticality", *cols], data, widths=[9, 36, 12, 11, 12, 13, 12, 11])
    for i in range(r0, r1 + 1):
        s.cell(i, 9, f"=SUM(D{i}:H{i})")
    s.cell(4, 9, "Total").fill, s.cell(4, 9).font = HDR_FILL, HDR_FONT

    s = wb.create_sheet("P2 Key activity (checklist)")
    _title(s, "Compliance by key activity · per site (checklist)", "Assessed items only (compliant / partial / non-compliant).")
    data = []
    for sid in order:
        acts = {}
        for r in by[sid]:
            if r["status_level"] in STATUS:
                acts.setdefault(r["activity"], Counter())[r["status_level"]] += 1
        for a, c in acts.items():
            data.append([sid, meta[sid]["site"], a, _en(tr, a), *[c.get(k, 0) for k in STATUS]])
    r0, r1 = _table(s, 4, 1, ["Site ID", "Site", "Key activity (original)", "Key activity (English)", *STATUS], data,
                    widths=[9, 30, 45, 45, 11, 12, 13])
    for i in range(r0, r1 + 1):
        s.cell(i, 8, f"=SUM(E{i}:G{i})")
    s.cell(4, 8, "Total").fill, s.cell(4, 8).font = HDR_FILL, HDR_FONT
    # one chart per site (as on page 2), stacked down column J
    pos, start = 4, r0
    for k in range(r0, r1 + 2):
        if k > r1 or s.cell(k, 1).value != s.cell(start, 1).value:
            n = k - start
            _stack_chart(s, f"{s.cell(start, 1).value} · {s.cell(start, 2).value}",
                         Reference(s, min_col=4, min_row=start, max_row=k - 1), 5, 3, start, k - 1, f"J{pos}",
                         [GREEN, ORANGE, RED], horizontal=True, height=max(6, 2 + 0.55 * n), width=20, titles=STATUS)
            pos += int(max(6, 2 + 0.55 * n) * 2) + 2
            start = k

    s = wb.create_sheet("P2 Report dots")
    _title(s, "Site-visit report · compliance dots per site")
    rep_cols = ["Compliant", "Partially compliant", "Non-compliant", "Not assessed"]
    data = []
    for i, sid in enumerate([k for k in order if k in reports]):
        c = Counter(k["status"] for k in reports[sid].get("key_activities", []))
        data.append([sid, meta[sid]["site"], *[c.get(k, 0) for k in rep_cols], f"=SUM(C{5 + i}:F{5 + i})"])
    r0, r1 = _table(s, 4, 1, ["Site ID", "Site", *rep_cols, "Key activities"], data, widths=[9, 36, 11, 12, 13, 11, 12])
    _stack_chart(s, "Report compliance dots per site", Reference(s, min_col=2, min_row=r0, max_row=r1), 3, 4, r0, r1,
                 "I4", [GREEN, ORANGE, RED, "D9DCE2"], horizontal=True, height=14)

    # ── Page 3 ──
    blocks = [(f"{cl} · {t}", deck["cluster_figures"][cl][t]) for cl in CLUSTERS for t in TYPES
              if t in deck["cluster_figures"][cl]]
    _crit_sheet(wb, "P3 Cluster business", "Compliance status per criticality of business · per cluster (deck)",
                "Global deck pp.14–19.", blocks)

    s = wb.create_sheet("P3 Cluster KPIs")
    _title(s, "Cluster KPI tiles (deck figures)")
    data = []
    for i, cl in enumerate(CLUSTERS):
        f = deck["cluster_figures"][cl]
        tot = [sum(f[t][c][k] for t in f for c in CRIT) for k in range(3)]
        hi = [sum(f[t]["High"][k] for t in f) for k in range(3)]
        n = 5 + i
        data.append([cl, len(f), *tot, f"=SUM(C{n}:E{n})", f"=C{n}/F{n}", *hi])
    r0, r1 = _table(s, 4, 1, ["Cluster", "Types of business", "Compliant", "Partially compliant", "Non-compliant",
                              "Requirements assessed", "% compliant", "High · compliant", "High · partial",
                              "High · non-compliant"], data, widths=[20, 10, 11, 12, 13, 12, 11, 11, 11, 12])
    for i in range(r0, r1 + 1):
        s.cell(i, 7).number_format = "0%"

    s = wb.create_sheet("P3 Cluster mix")
    _title(s, "Compliance mix by type of business (all criticality levels, deck figures)")
    data = []
    for cl in CLUSTERS:
        f = deck["cluster_figures"][cl]
        for t in TYPES:
            if t in f:
                data.append([cl, t, *[sum(f[t][c][k] for c in CRIT) for k in range(3)]])
    r0, r1 = _table(s, 4, 1, ["Cluster", "Type of business", *STATUS, "Total", "% compliant", "% partial", "% non-compliant"],
                    data, widths=[20, 14, 11, 12, 13, 9, 11, 10, 13])
    for i in range(r0, r1 + 1):
        s.cell(i, 6, f"=SUM(C{i}:E{i})")
        for j, col in enumerate("CDE"):
            c = s.cell(i, 7 + j, f"={col}{i}/F{i}")
            c.number_format = "0%"
    _stack_chart(s, "Compliance mix by type of business (100%)", Reference(s, min_col=1, max_col=2, min_row=r0, max_row=r1),
                 3, 3, r0, r1, "K4", [GREEN, ORANGE, RED], horizontal=True, pct=True, height=12, y_title="Share")

    s = wb.create_sheet("P3 Site heatmap")
    _title(s, "Share compliant by site and criticality (as charted)", "Deck site figures; Mexico in-house services = "
                                                                      "deck aggregate of Iberdrola, Schneider and Braskem.")
    data = []
    for cl in CLUSTERS:
        used = set()
        for sid in [k for k in order if meta[k]["cluster"] == cl]:
            if sid in deck["site_figures"]:
                f, lab = deck["site_figures"][sid], meta[sid]["site"]
            elif sid in deck["group_figures"]["MX-SERVICES"]["sites"] and "g" not in used:
                used.add("g"); f, lab = deck["group_figures"]["MX-SERVICES"]["figures"], "Mexico in-house services"
            else:
                continue
            data.append([cl, lab] + [x for c in CRIT for x in (f[c][0], sum(f[c]))])
    hdr = ["Cluster", "Site"] + [x for c in CRIT for x in (f"{c} · compliant", f"{c} · total")]
    r0, r1 = _table(s, 4, 1, hdr + [f"{c} · % compliant" for c in CRIT], [d + ["", "", ""] for d in data],
                    widths=[20, 34, 11, 9, 11, 9, 11, 9, 12, 12, 12])
    for i in range(r0, r1 + 1):
        for j, (a, b) in enumerate([("C", "D"), ("E", "F"), ("G", "H")]):
            c = s.cell(i, 9 + j, f'=IF({b}{i}=0,"–",{a}{i}/{b}{i})')
            c.number_format = "0%"
    s.conditional_formatting.add(f"I{r0}:K{r1}", ColorScaleRule(start_type="num", start_value=0, start_color=RED,
                                                                 mid_type="num", mid_value=0.5, mid_color="F6D48A",
                                                                 end_type="num", end_value=1, end_color=GREEN))

    # ── Page 4 ──
    s = wb.create_sheet("P4 Global high")
    _title(s, "Compliance status by HIGH criticality per type of business in the different clusters", "Global deck pp.12–13.")
    data = [[r["business_type"], r["cluster"], " + ".join(r["flags"]), *r["values"]] for r in deck["global_high"]]
    r0, r1 = _table(s, 4, 1, ["Type of business", "Cluster", "Countries", *STATUS, "Total"],
                    [d + [f"=SUM(D{5 + i}:F{5 + i})"] for i, d in enumerate(data)], widths=[14, 20, 10, 11, 12, 13, 9])
    s.cell(r1 + 1, 3, "Total").font = Font(bold=True)
    for j, col in enumerate("DEFG"):
        s.cell(r1 + 1, 4 + j, f"=SUM({col}{r0}:{col}{r1})").font = Font(bold=True)
    gh_tot = r1 + 1
    for i in range(r0, r1 + 1):
        s.cell(i, 9, f'=A{i}&" · "&C{i}')
    s.cell(4, 9, "Chart label").fill, s.cell(4, 9).font = HDR_FILL, HDR_FONT
    _stack_chart(s, "High criticality by type of business and cluster", Reference(s, min_col=9, min_row=r0, max_row=r1),
                 4, 3, r0, r1, "K4", [GREEN, ORANGE, RED], height=12, width=26)

    s = wb.create_sheet("P4 High by cluster")
    _title(s, "High criticality by type of business · one chart per cluster (deck)")
    top = 4
    for bi, cl in enumerate(CLUSTERS):
        f = deck["cluster_figures"][cl]
        s.cell(top, 1, cl).font = Font(bold=True, color=NAVY)
        data = [[t, *f[t]["High"], f"=SUM(B{top + 2 + i}:D{top + 2 + i})"] for i, t in enumerate([t for t in TYPES if t in f])]
        r0, r1 = _table(s, top + 1, 1, ["Type of business", *STATUS, "Total"], data, widths=[16, 12, 18, 14, 9])
        _stack_chart(s, f"{cl} · high criticality", Reference(s, min_col=1, min_row=r0, max_row=r1), 2, 3, r0, r1,
                     f"G{top}", [GREEN, ORANGE, RED], width=15, height=7.5)
        top = max(r1 + 3, top + 17)

    s = wb.create_sheet("P4 High mix")
    _title(s, "High-criticality compliance mix · per cluster and per type of business")
    data = [[cl, *[sum(deck["cluster_figures"][cl][t]["High"][k] for t in deck["cluster_figures"][cl]) for k in range(3)]]
            for cl in CLUSTERS]
    r0, r1 = _table(s, 4, 1, ["Cluster", *STATUS, "Total", "% compliant", "% partial", "% non-compliant"],
                    [d + ["", "", "", ""] for d in data], widths=[20, 11, 12, 13, 9, 11, 10, 13])
    for i in range(r0, r1 + 1):
        s.cell(i, 5, f"=SUM(B{i}:D{i})")
        for j, col in enumerate("BCD"):
            s.cell(i, 6 + j, f"={col}{i}/E{i}").number_format = "0%"
    _stack_chart(s, "High criticality mix per cluster", Reference(s, min_col=1, min_row=r0, max_row=r1), 2, 3, r0, r1,
                 "K4", [GREEN, ORANGE, RED], horizontal=True, pct=True, y_title="%", height=7.5)
    t0 = r1 + 4
    data = [[t, *[sum(r["values"][k] for r in deck["global_high"] if r["business_type"] == t) for k in range(3)]] for t in TYPES]
    r0, r1 = _table(s, t0, 1, ["Type of business", *STATUS, "Total", "% compliant", "% partial", "% non-compliant"],
                    [d + ["", "", "", ""] for d in data])
    for i in range(r0, r1 + 1):
        s.cell(i, 5, f"=SUM(B{i}:D{i})")
        for j, col in enumerate("BCD"):
            s.cell(i, 6 + j, f"={col}{i}/E{i}").number_format = "0%"
    _stack_chart(s, "High criticality mix per type of business", Reference(s, min_col=1, min_row=r0, max_row=r1), 2, 3,
                 r0, r1, f"K{t0}", [GREEN, ORANGE, RED], horizontal=True, pct=True, y_title="%", height=7.5)

    allc = [(cl, {c: [sum(deck["cluster_figures"][cl][t][c][k] for t in deck["cluster_figures"][cl]) for k in range(3)]
                  for c in CRIT}) for cl in CLUSTERS]
    _crit_sheet(wb, "P4 All criticality", "All criticality levels · clusters side by side (deck)",
                "Sum over types of business of the cluster charts (global deck pp.14–19).", allc)

    s = wb.create_sheet("P4 KPIs")
    _title(s, "Page 4 headline numbers (high criticality)", "Must match the executive summary: 184 compliant, "
                                                            "161 critical gaps (89 non-compliant + 72 partially compliant).")
    _table(s, 4, 1, ["Measure", "Value"], [
        ("High-criticality items", f"='P4 Global high'!G{gh_tot}"), ("Compliant", f"='P4 Global high'!D{gh_tot}"),
        ("Partially compliant", f"='P4 Global high'!E{gh_tot}"), ("Non-compliant", f"='P4 Global high'!F{gh_tot}"),
        ("Critical gaps (partial + non-compliant)", "=B7+B8")], widths=[40, 12])

    # ── QA ──
    s = wb.create_sheet("QA Deck vs checklist")
    _title(s, "Deck chart figures vs counts from the checklists",
           "The dashboard shows deck figures (as presented to the client). FALSE = the deck and the checklist differ.")
    dv = qa.get("deck_vs_checklist", [])
    data = [[d["id"], d["criticality"], *d["deck"], *d["checklist"], d["match"]] for d in dv]
    _table(s, 4, 1, ["Site", "Criticality", "Deck · C", "Deck · P", "Deck · NC", "Checklist · C", "Checklist · P",
                     "Checklist · NC", "Match"], data, widths=[22, 11, 9, 9, 9, 12, 12, 13, 8])
    s = wb.create_sheet("QA Deck internal")
    _title(s, "Deck charts that differ from the sum of their own site charts", "Values as printed in the decks.")
    data = [[d["chart"], d["sites"], str(d["deck"]), str(d["sum_of_site_slides"])] for d in qa.get("deck_internal_inconsistencies", [])]
    _table(s, 4, 1, ["Chart", "Site slides", "Deck shows [C, P, NC]", "Sum of site slides [C, P, NC]"], data,
           widths=[60, 18, 20, 26])
    s = wb.create_sheet("QA Reconciliation")
    _title(s, "Source Excel → dashboard dataset reconciliation")
    data = [[x["site_id"], x["file"], x["source_rows_with_id"], x["dataset_rows"], x["row_count_match"]]
            for x in qa.get("reconciliation_summary", [])]
    r0, r1 = _table(s, 4, 1, ["Site", "Source file", "Rows in source", "Rows in dashboard", "Match"], data,
                    widths=[9, 70, 13, 15, 8])
    s.cell(r1 + 1, 2, "Total").font = Font(bold=True)
    s.cell(r1 + 1, 3, f"=SUM(C{r0}:C{r1})").font = Font(bold=True)
    s.cell(r1 + 1, 4, f"=SUM(D{r0}:D{r1})").font = Font(bold=True)
    s.cell(r1 + 3, 1, f"Cell-by-cell comparison of every field: {qa.get('cell_mismatches', 'n/a')} mismatches "
                      "(status values read from cell colour are listed on the records workbook 'Status mapping' sheet).").font = NOTE_FONT
    return _bytes(wb)
